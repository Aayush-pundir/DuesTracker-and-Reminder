"""Excel importers. Re-uploading a fresh Consol file is safe (idempotent *snapshot* sync):
 - the sheet says what is outstanding NOW; invoices are created/updated so outstanding == sheet
 - invoices that disappeared from the sheet are closed with a SNAPSHOT adjustment (visible in the audit trail)
 - payments already recorded in the app are preserved (amount = sheet pending + paid-in-app)."""
import calendar, datetime, re
import openpyxl
from .db import audit, get_settings
from . import ledger

def _month(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return datetime.date(v.year, v.month, 1).isoformat()
    if isinstance(v, str):
        for fmt in ("%b %Y", "%B %Y", "%Y-%m-%d"):
            try:
                d = datetime.datetime.strptime(v.strip(), fmt); return datetime.date(d.year, d.month, 1).isoformat()
            except ValueError:
                pass
    return None

def _due(period, credit_days):
    y, m = int(period[:4]), int(period[5:7])
    return (datetime.date(y, m, calendar.monthrange(y, m)[1]) + datetime.timedelta(days=int(credit_days))).isoformat()

def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0

def _key(name):
    return re.sub(r"\s+", " ", str(name or "")).strip().lower()

def _customer(conn, cache, name):
    """Find-or-create by whitespace/case-insensitive name. `cache` is shared for the whole import run."""
    if not cache.get("_loaded"):
        cache["_loaded"] = True
        for r in conn.execute("SELECT id,name FROM customers"):
            cache[_key(r["name"])] = r["id"]
    k = _key(name)
    if k not in cache:
        cache[k] = conn.execute("INSERT INTO customers(name) VALUES(?)", (re.sub(r"\s+", " ", str(name)).strip(),)).lastrowid
    return cache[k]


def _tag_reseller(conn, cid):
    c = conn.execute("SELECT name,email_to,reseller_id FROM customers WHERE id=?", (cid,)).fetchone()
    if c["reseller_id"]:
        return
    for rs in conn.execute("SELECT * FROM resellers"):
        doms = [d.strip().lower() for d in rs["email_domains"].split(",") if d.strip()]
        marks = [m.strip().lower() for m in rs["name_markers"].split(",") if m.strip()]
        if any(d in (c["email_to"] or "").lower() for d in doms) or any(m in c["name"].lower() for m in marks):
            conn.execute("UPDATE customers SET reseller_id=? WHERE id=?", (rs["id"], cid)); return

def _sync_invoices(conn, wanted, kind, credit_days):
    """wanted: {(cid, period, invoice_no): (pending, description)} -> make outstanding match."""
    seen, st = set(), {"created": 0, "updated": 0, "closed_by_snapshot": 0}
    for (cid, period, no), (pending, desc) in wanted.items():
        row = conn.execute("SELECT id FROM invoices WHERE customer_id=? AND kind=? AND period=? AND invoice_no=?",
                           (cid, kind, period, no)).fetchone()
        if row:
            paid = conn.execute("SELECT COALESCE(SUM(amount),0) FROM allocations WHERE invoice_id=?", (row["id"],)).fetchone()[0]
            conn.execute("UPDATE invoices SET amount=?, description=?, status='OPEN' WHERE id=?", (pending + paid, desc, row["id"]))
            ledger.refresh_invoice(conn, row["id"]); iid = row["id"]; st["updated"] += 1
        else:
            iid = conn.execute("INSERT INTO invoices(customer_id,kind,period,invoice_no,description,amount,due_date) VALUES(?,?,?,?,?,?,?)",
                               (cid, kind, period, no, desc, pending, _due(period, credit_days))).lastrowid
            st["created"] += 1
        seen.add(iid)
    for r in conn.execute("""SELECT i.id,i.customer_id, i.amount-COALESCE((SELECT SUM(amount) FROM allocations WHERE invoice_id=i.id),0) bal
                             FROM invoices i WHERE kind=? AND status IN ('OPEN','PARTIAL')""", (kind,)).fetchall():
        if r["id"] not in seen and r["bal"] > ledger.EPS:
            pid = conn.execute("INSERT INTO payments(customer_id,mode,amount,received_on,narration) VALUES(?,?,?,?,?)",
                               (r["customer_id"], "SNAPSHOT", r["bal"], ledger.today(), "Absent from latest Consol sheet")).lastrowid
            conn.execute("INSERT INTO allocations(payment_id,invoice_id,amount) VALUES(?,?,?)", (pid, r["id"], r["bal"]))
            ledger.refresh_invoice(conn, r["id"]); st["closed_by_snapshot"] += 1
    return st

def import_consol(conn, path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    s = get_settings(conn); cd = s["credit_days"]; cache = {}
    report = {}
    # 1) Spend-management dues: customer x month matrix
    ws = wb["Summary SM"]; rows = list(ws.iter_rows(min_row=2, values_only=True))
    months = {i: _month(v) for i, v in enumerate(rows[0]) if i >= 3 and _month(v)}
    wanted = {}
    for r in rows[1:]:
        name = r[0]
        if not name or str(name).strip().lower() == "grand total":
            continue
        cid = _customer(conn, cache, name)
        kam = (r[2] or "").strip()
        conn.execute("UPDATE customers SET enterprise_id=COALESCE(?,enterprise_id), kam=CASE WHEN ?!='' THEN ? ELSE kam END WHERE id=?",
                     (r[1], kam.title() if kam else "", kam.title() if kam else "", cid))
        if "freeze" in kam.lower() or "stopped" in kam.lower():
            conn.execute("UPDATE customers SET status='FROZEN', remind=0 WHERE id=?", (cid,))
        for i, p in months.items():
            v = _num(r[i]) if i < len(r) else 0
            if v > 0.5:
                wanted[(cid, p, "")] = (v, "Spend management")
    report["spend_mgmt"] = _sync_invoices(conn, wanted, "SPEND_MGMT", cd)
    # 2) Everything that is not spend management: invoice level
    ws = wb["Other than S Mgt"]; wanted = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not (r[1] and isinstance(r[3], (int, float)) and r[4]):
            continue
        cid = _customer(conn, cache, r[1]); p = _month(r[5]) or ledger.today()[:8] + "01"
        wanted[(cid, p, str(r[4]).strip())] = (float(r[3]), str(r[6] or "")[:200])
        if r[7]:
            conn.execute("UPDATE customers SET kam=COALESCE(NULLIF(kam,''),?) WHERE id=?", (str(r[7]), cid))
    report["other"] = _sync_invoices(conn, wanted, "OTHER", cd)
    # 3) Ops view: how we can collect from each customer + billing contacts
    if "Spend Management" in wb.sheetnames:
        ws = wb["Spend Management"]; seen = set()
        for r in ws.iter_rows(min_row=3, max_col=11, values_only=True):
            if not r[1] or str(r[1]).strip().lower() in ("grand total", "total"):
                continue
            k = _key(r[1])
            if k in seen:
                continue
            seen.add(k)
            cid = _customer(conn, cache, r[1]); ops = str(r[5] or "").strip(); o = ops.lower()
            mode = "WALLET" if "running balance" in o else None
            conn.execute("UPDATE customers SET ops_status=?, auto_debit_allowed=?, collection_mode=COALESCE(?,collection_mode),"
                         " email_to=CASE WHEN email_to='' THEN ? ELSE email_to END, enterprise_id=COALESCE(enterprise_id,?) WHERE id=?",
                         (ops, 0 if o.startswith("do not auto debit") else 1, mode, str(r[10] or ""), r[9], cid))
    # 4) Reseller tagging (IDFC / Stride) from names + e-mail domains
    for k, cid in cache.items():
        if k != "_loaded":
            _tag_reseller(conn, cid)
    if "Stride Billing" in wb.sheetnames:
        sid = conn.execute("SELECT id FROM resellers WHERE name='Stride One'").fetchone()["id"]
        for r in wb["Stride Billing"].iter_rows(min_row=2, max_col=1, values_only=True):
            if r[0] and _key(r[0]) in cache:
                conn.execute("UPDATE customers SET reseller_id=COALESCE(reseller_id,?) WHERE id=?", (sid, cache[_key(r[0])]))
    audit(conn, "import_consol", str(report)); conn.commit()
    report["customers"] = conn.execute("SELECT COUNT(*) FROM customers").fetchone()[0]
    return report

def import_contacts(conn, path):
    """DuesReminder workbook (sheet DATA): customer, email-to, email-cc  -> fills/updates contacts."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True, keep_vba=False)
    n, cache = 0, {}
    for r in wb["DATA"].iter_rows(min_row=3, max_col=3, values_only=True):
        if not r[0]:
            continue
        cid = _customer(conn, cache, r[0])
        conn.execute("UPDATE customers SET email_to=COALESCE(NULLIF(?,''),email_to), email_cc=COALESCE(NULLIF(?,''),email_cc) WHERE id=?",
                     (r[1] or "", r[2] or "", cid))
        _tag_reseller(conn, cid); n += 1
    conn.commit(); audit(conn, "import_contacts", str(n))
    return {"contacts_updated": n}

def read_bank_statement(path):
    """Generic bank-statement reader: finds header row containing date / narration / credit columns (CSV or XLSX)."""
    import csv
    if path.lower().endswith(".csv"):
        data = list(csv.reader(open(path, newline="", encoding="utf-8-sig")))
    else:
        data = [list(r) for r in openpyxl.load_workbook(path, data_only=True, read_only=True).worksheets[0].iter_rows(values_only=True)]
    hi = next(i for i, r in enumerate(data) if sum(1 for c in r if c and re.search(r"date|narration|description|particulars|credit|deposit|ref|chq", str(c), re.I)) >= 3)
    h = [str(c or "").lower() for c in data[hi]]
    def col(*pats):
        return next((i for i, x in enumerate(h) if any(re.search(p, x) for p in pats)), None)
    d, n, cr, rf = col(r"date"), col(r"narration|description|particulars|remarks"), col(r"credit|deposit|cr amount"), col(r"ref|chq|utr")
    out = []
    for r in data[hi + 1:]:
        if not r or r[d] in (None, ""):
            continue
        dv = r[d]
        if isinstance(dv, str):
            for f in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y", "%d-%b-%Y"):
                try: dv = datetime.datetime.strptime(dv.strip(), f); break
                except ValueError: pass
        amt = str(r[cr] or "0").replace(",", "")
        out.append({"date": dv.date().isoformat() if hasattr(dv, "date") else str(dv), "amount": _num(amt),
                    "narration": str(r[n] or ""), "reference": str(r[rf] or "") if rf is not None else ""})
    return out
