"""Management reports. The dues matrix reproduces the familiar customer x month view of the old Consol sheet (live, always current)."""
import datetime, io
import openpyxl
from openpyxl.styles import Font, PatternFill
from . import ledger

BAL = ledger.BAL
BUCKETS = ["Not yet due", "1-30 d", "31-60 d", "61-90 d", "90+ d"]

def _bucket(days):
    return 0 if days <= 0 else 1 if days <= 30 else 2 if days <= 60 else 3 if days <= 90 else 4

def matrix(conn, kam="", payer=""):
    rows = conn.execute(f"""SELECT c.id, c.name, c.kam, substr(i.period,1,7) m, SUM({BAL}) bal FROM invoices i JOIN customers c ON c.id=i.customer_id
        WHERE i.status IN ('OPEN','PARTIAL') AND (?='' OR c.kam=?) AND (?='' OR i.payer=?) GROUP BY c.id, m HAVING bal>0.5""", (kam, kam, payer, payer)).fetchall()
    months = sorted({r["m"] for r in rows}); by = {}
    for r in rows:
        d = by.setdefault(r["id"], {"id": r["id"], "name": r["name"], "kam": r["kam"], "cells": {}, "total": 0})
        d["cells"][r["m"]] = round(r["bal"]); d["total"] += round(r["bal"])
    out = sorted(by.values(), key=lambda x: -x["total"])
    return {"months": months, "rows": out, "totals": {m: sum(r["cells"].get(m, 0) for r in out) for m in months}, "grand": sum(r["total"] for r in out)}

def ageing(conn, kam=""):
    t = datetime.date.today(); by = {}
    for r in conn.execute(f"""SELECT c.id, c.name, c.kam, i.due_date, {BAL} bal FROM invoices i JOIN customers c ON c.id=i.customer_id
                              WHERE i.status IN ('OPEN','PARTIAL') AND i.hold=0 AND (?='' OR c.kam=?)""", (kam, kam)):
        if r["bal"] < 0.5: continue
        d = by.setdefault(r["id"], {"id": r["id"], "name": r["name"], "kam": r["kam"], "b": [0] * 5})
        d["b"][_bucket((t - datetime.date.fromisoformat(r["due_date"])).days)] += r["bal"]
    out = [{**d, "b": [round(x) for x in d["b"]], "total": round(sum(d["b"]))} for d in by.values()]
    out.sort(key=lambda x: -x["total"])
    return {"buckets": BUCKETS, "rows": out, "totals": [sum(r["b"][i] for r in out) for i in range(5)], "grand": sum(r["total"] for r in out)}

def trend(conn, n=12):
    t = datetime.date.today().replace(day=1); ms = []
    for _ in range(n):
        ms.append(t.strftime("%Y-%m")); t = (t - datetime.timedelta(days=1)).replace(day=1)
    ms.reverse()
    billed = dict(conn.execute("SELECT substr(invoice_date,1,7), SUM(amount) FROM invoices WHERE status!='CANCELLED' GROUP BY 1").fetchall())
    coll = dict(conn.execute("SELECT substr(received_on,1,7), SUM(amount+tds_amount) FROM payments WHERE status='ALLOCATED' AND mode!='SNAPSHOT' GROUP BY 1").fetchall())
    return [{"month": m, "billed": round(billed.get(m) or 0), "collected": round(coll.get(m) or 0)} for m in ms]

def kam_scorecard(conn):
    t = datetime.date.today(); iso = t.isoformat(); month = t.replace(day=1).isoformat(); out = {}
    for r in conn.execute(f"""SELECT COALESCE(NULLIF(c.kam,''),'-') k, c.id, i.due_date, {BAL} bal FROM invoices i JOIN customers c ON c.id=i.customer_id
                              WHERE i.status IN ('OPEN','PARTIAL')"""):
        d = out.setdefault(r["k"], {"kam": r["k"], "customers": set(), "outstanding": 0, "overdue": 0, "collected": 0, "broken_promises": 0})
        d["customers"].add(r["id"]); d["outstanding"] += r["bal"]; d["overdue"] += r["bal"] if r["due_date"] < iso else 0
    for r in conn.execute("""SELECT COALESCE(NULLIF(c.kam,''),'-') k, SUM(p.amount+p.tds_amount) amt FROM payments p JOIN customers c ON c.id=p.customer_id
                             WHERE p.received_on>=? AND p.status='ALLOCATED' AND p.mode!='SNAPSHOT' GROUP BY k""", (month,)):
        out.setdefault(r["k"], {"kam": r["k"], "customers": set(), "outstanding": 0, "overdue": 0, "collected": 0, "broken_promises": 0})["collected"] = r["amt"]
    for r in conn.execute("SELECT COALESCE(NULLIF(kam,''),'-') k, COUNT(*) n FROM customers WHERE promise_date<? GROUP BY k", (iso,)):
        if r["k"] in out: out[r["k"]]["broken_promises"] = r["n"]
    res = [{**d, "customers": len(d["customers"]), "outstanding": round(d["outstanding"]), "overdue": round(d["overdue"]), "collected": round(d["collected"])} for d in out.values()]
    return sorted(res, key=lambda x: -x["outstanding"])

def _xlsx(title, header, rows, total_row=None):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = title[:30]
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="4F46E5")
    for r in rows: ws.append(r)
    if total_row: ws.append(total_row); [setattr(c, "font", Font(bold=True)) for c in ws[ws.max_row]]
    ws.freeze_panes = "C2"
    for i, col in enumerate(ws.columns, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 38 if i == 1 else 16
    b = io.BytesIO(); wb.save(b); return b.getvalue()

def matrix_xlsx(conn, kam=""):
    m = matrix(conn, kam)
    return _xlsx("Dues matrix", ["Customer", "KAM", *m["months"], "Total"],
                 [[r["name"], r["kam"], *[r["cells"].get(x, None) for x in m["months"]], r["total"]] for r in m["rows"]],
                 ["Grand total", "", *[m["totals"][x] for x in m["months"]], m["grand"]])

def ageing_xlsx(conn, kam=""):
    a = ageing(conn, kam)
    return _xlsx("Ageing", ["Customer", "KAM", *a["buckets"], "Total"], [[r["name"], r["kam"], *r["b"], r["total"]] for r in a["rows"]],
                 ["Grand total", "", *a["totals"], a["grand"]])

def all_data_xlsx(conn):
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    for name, sql in (("customers", "SELECT * FROM customers"), ("invoices", f"SELECT i.*, {BAL} balance FROM invoices i"),
                      ("payments", "SELECT * FROM payments"), ("allocations", "SELECT * FROM allocations"), ("notes", "SELECT * FROM notes")):
        cur = conn.execute(sql); ws = wb.create_sheet(name); ws.append([d[0] for d in cur.description])
        for r in cur: ws.append(list(r))
    b = io.BytesIO(); wb.save(b); return b.getvalue()
